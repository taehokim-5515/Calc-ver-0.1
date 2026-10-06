import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
from pathlib import Path
import calendar
import shutil
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io

# 페이지 설정
st.set_page_config(
    page_title="재고 관리 시스템",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─────────────────────────────────────────────────────────────
# 발주 정보 입력 규격
# ─────────────────────────────────────────────────────────────
ORDER_INFO_COLUMNS = ['부가세 구분', '제품 구분', '업체명', '단량(KG)', '단가(원)', '담당자 이메일']
VAT_OPTIONS = ['과세', '면세']
PRODUCT_TYPE_OPTIONS = ['냉동', '냉장', '포장', '액상', '수제']
WEEKDAYS_KR = ['월', '화', '수', '목', '금', '토', '일']

# ─────────────────────────────────────────────────────────────
# 거래처 DB (앱과 분리된 CSV 파일, 업로드한 파일로 전체 교체)
# ─────────────────────────────────────────────────────────────
VENDOR_DB_FILE = Path(__file__).parent / '거래처DB.csv'
VENDOR_DB_BACKUP_FILE = Path(__file__).parent / '거래처DB_backup.csv'
VENDOR_NAME_COLUMNS = ['업체명', '거래처명', '거래처']
VENDOR_EMAIL_COLUMNS = ['담당자 이메일', '담당자이메일', '이메일', 'E-mail', 'email']

def normalize_code(value):
    """원료코드를 비교 가능한 문자열로 정규화 (Excel에서 1010101.0 처럼 읽히는 경우 대비)"""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()

def normalize_vendor_df(raw):
    """업로드/저장된 거래처 데이터를 '업체명', '담당자 이메일' 2개 열로 정리하는 함수"""
    raw = raw.copy()
    raw.columns = [str(c).strip() for c in raw.columns]

    vendor_col = next((c for c in VENDOR_NAME_COLUMNS if c in raw.columns), None)
    if vendor_col is None:
        raise ValueError(f"업체명 열을 찾을 수 없습니다. (인식 가능한 열 이름: {', '.join(VENDOR_NAME_COLUMNS)})")
    email_col = next((c for c in VENDOR_EMAIL_COLUMNS if c in raw.columns), None)

    vdb = pd.DataFrame({
        '업체명': raw[vendor_col],
        '담당자 이메일': raw[email_col] if email_col else '',
    }).fillna('')
    vdb['업체명'] = vdb['업체명'].astype(str).str.strip()
    vdb['담당자 이메일'] = vdb['담당자 이메일'].astype(str).str.strip().replace('미입력', '')
    vdb = vdb[vdb['업체명'] != '']

    # 같은 업체가 여러 번 나오면 이메일이 있는 행을 우선
    vdb = (vdb.assign(_no_email=vdb['담당자 이메일'] == '')
              .sort_values('_no_email', kind='stable')
              .drop_duplicates(subset=['업체명'], keep='first')
              .sort_index()
              .drop(columns='_no_email')
              .reset_index(drop=True))
    return vdb

def read_table_file(file_or_path, name):
    """CSV/Excel 파일을 문자열 기준으로 읽는 함수 (CSV는 UTF-8 → CP949 순으로 시도)"""
    if str(name).lower().endswith('.csv'):
        try:
            return pd.read_csv(file_or_path, dtype=str, encoding='utf-8-sig')
        except UnicodeDecodeError:
            if hasattr(file_or_path, 'seek'):
                file_or_path.seek(0)
            return pd.read_csv(file_or_path, dtype=str, encoding='cp949')
    return pd.read_excel(file_or_path, dtype=str)

def load_vendor_db():
    """저장된 거래처 DB를 읽어오는 함수 (없으면 빈 DB)"""
    empty = pd.DataFrame(columns=['업체명', '담당자 이메일'])
    if not VENDOR_DB_FILE.exists():
        return empty
    try:
        return normalize_vendor_df(read_table_file(VENDOR_DB_FILE, VENDOR_DB_FILE.name))
    except Exception as e:
        st.warning(f"⚠️ 거래처 DB 파일을 읽지 못했습니다: {str(e)}")
        return empty

def save_vendor_db(vdb):
    """거래처 DB를 저장하는 함수 (기존 파일은 백업 후 교체)"""
    if VENDOR_DB_FILE.exists():
        shutil.copyfile(VENDOR_DB_FILE, VENDOR_DB_BACKUP_FILE)
    vdb.to_csv(VENDOR_DB_FILE, index=False, encoding='utf-8-sig')

def is_weekend(date):
    """주말인지 확인하는 함수 (토요일: 5, 일요일: 6)"""
    return date.weekday() >= 5

def get_next_monday(date):
    """다음 월요일 날짜를 반환하는 함수"""
    days_until_monday = (7 - date.weekday()) % 7
    if days_until_monday == 0:
        days_until_monday = 7
    return date + timedelta(days=days_until_monday)

def parse_date(date_str):
    """날짜 문자열을 datetime 객체로 변환"""
    if isinstance(date_str, datetime):
        return date_str.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(date_str, str):
        return datetime.strptime(date_str.strip().split()[0], '%Y-%m-%d')
    return None

def is_blank(value):
    """빈 값 여부 확인 (None, NaN, 빈 문자열)"""
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ''
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False

def parse_scheduled_receipts(edited_df, label_to_code, today):
    """입고 예정 입력표를 {원료코드: {날짜: 입고량}} 형태로 변환하는 함수
    반환: (입고 예정 dict, 과거 날짜로 제외된 건수, 입력이 덜 된 건수)
    """
    receipts = {}
    past_count = 0
    incomplete_count = 0

    for _, r in edited_df.iterrows():
        label, receive_date, qty = r.get('원료'), r.get('입고일'), r.get('입고량(KG)')
        blanks = [is_blank(v) for v in (label, receive_date, qty)]
        if all(blanks):
            continue
        try:
            qty = float(qty)
        except (TypeError, ValueError):
            qty = 0
        if any(blanks) or qty <= 0 or label not in label_to_code:
            incomplete_count += 1
            continue

        receive_date = pd.to_datetime(receive_date).date()
        if receive_date < today:
            past_count += 1
            continue

        code = label_to_code[label]
        receipts.setdefault(code, {})
        receipts[code][receive_date] = receipts[code].get(receive_date, 0.0) + qty

    return receipts, past_count, incomplete_count

def collect_calendar_dates(results_data, receipt_events):
    """캘린더 표시용 날짜별 발주/예정입고 품목 목록을 만드는 함수"""
    order_dates = {}    # {날짜: [품목명]} - 발주 필요(자동 산출)
    receipt_dates = {}  # {날짜: [품목명]} - 입고 예정(사용자 입력)

    for item in results_data:
        if item['입고 필요일'] != '입고 필요 없음':
            for date in item['입고 필요일'].split(', '):
                order_dates.setdefault(date, []).append(item['품목명'])

    for event in receipt_events or []:
        receipt_dates.setdefault(event['날짜'], []).append(event['품목명'])

    all_dates = []
    for date in set(order_dates) | set(receipt_dates):
        try:
            all_dates.append(datetime.strptime(date, '%Y-%m-%d'))
        except ValueError:
            continue

    return order_dates, receipt_dates, all_dates

def display_calendar(results_data, receipt_events=None):
    """스트림릿에서 캘린더 형식으로 발주 필요일과 입고 예정일을 표시하는 함수"""
    try:
        order_dates, receipt_dates, all_dates = collect_calendar_dates(results_data, receipt_events)

        if not all_dates:
            st.info("📅 입고 일정이 없습니다.")
            return

        start_date = min(all_dates)
        end_date = max(all_dates)

        st.info(f"📅 입고 일정 기간: {start_date.strftime('%Y-%m-%d')} ~ {end_date.strftime('%Y-%m-%d')}")

        current_date = start_date.replace(day=1)
        end_month = end_date.replace(day=1)

        while current_date <= end_month:
            st.markdown(f"### 📅 {current_date.strftime('%Y년 %m월')}")

            col_headers = st.columns(7)
            for i, day in enumerate(WEEKDAYS_KR):
                with col_headers[i]:
                    st.markdown(f"**{day}**")

            cal = calendar.monthcalendar(current_date.year, current_date.month)

            for week in cal:
                cols = st.columns(7)
                for day_idx, day in enumerate(week):
                    with cols[day_idx]:
                        if day == 0:
                            st.markdown("<div style='height: 80px;'></div>", unsafe_allow_html=True)
                            continue

                        date_str = f"{current_date.year}-{current_date.month:02d}-{day:02d}"
                        is_weekend_day = day_idx >= 5
                        has_order = date_str in order_dates
                        has_receipt = date_str in receipt_dates

                        if has_order:
                            bg_color, border_color, text_color = "#FFE6E6", "#FF4444", "#000000"
                        elif has_receipt:
                            bg_color, border_color, text_color = "#E6F0FF", "#3B82F6", "#000000"
                        elif is_weekend_day:
                            bg_color, border_color, text_color = "#F0F0F0", "#CCCCCC", "#666666"
                        else:
                            bg_color, border_color, text_color = "#FFFFFF", "#DDDDDD", "#000000"

                        cell_content = f"""
                        <div style='
                            background-color: {bg_color};
                            border: 2px solid {border_color};
                            border-radius: 5px;
                            padding: 5px;
                            margin: 2px;
                            min-height: 70px;
                            color: {text_color};
                            font-size: 12px;
                        '>
                            <div style='font-weight: bold; margin-bottom: 3px;'>{day}</div>
                        """

                        tag_style = "color: white; padding: 1px 3px; margin: 1px 0; border-radius: 3px; font-size: 10px;"
                        for item_name in order_dates.get(date_str, []):
                            display_name = item_name if len(item_name) <= 8 else item_name[:8] + "..."
                            cell_content += f"<div style='background-color: #FF6666; {tag_style}'>{display_name}</div>"
                        for item_name in receipt_dates.get(date_str, []):
                            display_name = item_name if len(item_name) <= 7 else item_name[:7] + "..."
                            cell_content += f"<div style='background-color: #3B82F6; {tag_style}'>📥 {display_name}</div>"

                        cell_content += "</div>"
                        st.markdown(cell_content, unsafe_allow_html=True)

            st.markdown("---")

            if current_date.month == 12:
                current_date = current_date.replace(year=current_date.year + 1, month=1)
            else:
                current_date = current_date.replace(month=current_date.month + 1)

        # 범례 표시
        st.markdown("### 📖 범례")
        legend_cols = st.columns(4)
        legends = [
            ("#FFE6E6", "#FF4444", "🔴 발주 필요 (자동 산출)"),
            ("#E6F0FF", "#3B82F6", "🔵 입고 예정 (직접 입력)"),
            ("#F0F0F0", "#CCCCCC", "⚪ 주말"),
            ("#FFFFFF", "#DDDDDD", "⚫ 평일"),
        ]
        for col, (bg, border, label) in zip(legend_cols, legends):
            with col:
                st.markdown(f"""
                <div style='background-color: {bg}; border: 2px solid {border}; padding: 10px; border-radius: 5px; text-align: center; color: #000000;'>
                    <strong>{label}</strong>
                </div>
                """, unsafe_allow_html=True)

        # 입고 일정 상세 정보
        st.markdown("### 📋 입고 일정 상세")

        for date_str in sorted(set(order_dates) | set(receipt_dates)):
            date_obj = datetime.strptime(date_str, '%Y-%m-%d')
            weekday = WEEKDAYS_KR[date_obj.weekday()]
            orders = order_dates.get(date_str, [])
            receipts = receipt_dates.get(date_str, [])

            with st.expander(f"📅 {date_str} ({weekday}) - 발주 필요 {len(orders)}개 / 입고 예정 {len(receipts)}개"):
                for item in orders:
                    st.write(f"🔴 발주 필요: {item}")
                for item in receipts:
                    st.write(f"🔵 입고 예정: {item}")

    except Exception as e:
        st.error(f"캘린더 표시 중 오류 발생: {str(e)}")
        st.info("입고 일정을 목록 형태로 표시합니다.")

        rows = []
        for item in results_data:
            if item['입고 필요일'] != '입고 필요 없음':
                for date in item['입고 필요일'].split(', '):
                    rows.append({'날짜': date, '구분': '발주 필요', '품목명': item['품목명']})
        for event in receipt_events or []:
            rows.append({'날짜': event['날짜'], '구분': '입고 예정', '품목명': event['품목명']})

        if rows:
            st.dataframe(pd.DataFrame(rows).sort_values('날짜'), width='stretch')

def calculate_purchase_date(df, web_exclude_dates=None, web_include_dates=None, scheduled_receipts=None):
    """안전재고 도달일과 입고 필요일을 계산하는 함수
    scheduled_receipts: {원료코드: {date: 입고량(KG)}} — 해당 일자 시작 시점에 재고로 반영
    반환: (품목별 결과 리스트, 반영된 입고 예정 목록)
    """
    try:
        results = []
        receipt_events = []
        scheduled_receipts = scheduled_receipts or {}

        # 웹에서 입력받은 날짜를 datetime 객체로 변환
        web_exclude_datetimes = []
        if web_exclude_dates:
            if isinstance(web_exclude_dates, (list, tuple)):
                web_exclude_datetimes = [datetime.combine(date, datetime.min.time()) for date in web_exclude_dates]
            else:
                web_exclude_datetimes = [datetime.combine(web_exclude_dates, datetime.min.time())]

        web_include_datetimes = []
        if web_include_dates:
            if isinstance(web_include_dates, (list, tuple)):
                web_include_datetimes = [datetime.combine(date, datetime.min.time()) for date in web_include_dates]
            else:
                web_include_datetimes = [datetime.combine(web_include_dates, datetime.min.time())]

        for _, row in df.iterrows():
            # 기본 데이터 추출
            item_code = row['원료코드명']
            item_name = row['원료명']
            current_inventory = float(row['현재 재고'])
            daily_usage = (float(row['최소 사용량']) + float(row['최대 사용량'])) / 2
            calculation_days = int(row['계산 일자'])
            safety_stock = float(row['안전 재고'])
            purchase_amount = float(row['1회 구매량'])

            # 시작일 설정
            start_date = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            end_date = start_date + timedelta(days=calculation_days)

            # Excel 파일의 제외 날짜 처리 (열이 없어도 동작)
            excel_exclude_dates = []
            if pd.notna(row.get('제외 날짜')):
                date_strings = str(row.get('제외 날짜')).split(',')
                excel_exclude_dates = [parse_date(date_str) for date_str in date_strings if parse_date(date_str)]

            # Excel 파일의 포함 날짜 처리 (열이 없어도 동작)
            excel_include_dates = []
            if pd.notna(row.get('포함 날짜')):
                date_strings = str(row.get('포함 날짜')).split(',')
                excel_include_dates = [parse_date(date_str) for date_str in date_strings if parse_date(date_str)]

            # 최종 제외/포함 날짜 결합 (웹 입력이 우선, Excel 데이터와 합침)
            final_exclude_dates = web_exclude_datetimes + excel_exclude_dates
            final_include_dates = web_include_datetimes + excel_include_dates

            # 이 품목의 입고 예정
            item_receipts = scheduled_receipts.get(normalize_code(item_code), {})
            out_of_range = sum(1 for d in item_receipts if not (start_date.date() <= d <= end_date.date()))
            applied_receipts = []
            review_notes = []

            # 안전재고 도달일 계산
            current_date = start_date
            remaining_inventory = current_inventory
            safety_stock_info = []

            while current_date <= end_date:
                day = current_date.date()

                # 입고 예정: 가동일 여부와 상관없이 그날 시작 시점에 재고 반영
                if day in item_receipts:
                    qty = item_receipts[day]
                    remaining_inventory += qty
                    applied_receipts.append(f"{day:%Y-%m-%d}({qty:,.0f}kg)")
                    receipt_events.append({
                        '날짜': f"{day:%Y-%m-%d}",
                        '품목코드': item_code,
                        '품목명': item_name,
                        '입고량(KG)': qty,
                    })

                # 주말이면서 포함 날짜에 없는 경우 스킵
                is_included_date = any(include_date and include_date.date() == day
                                    for include_date in final_include_dates)
                if is_weekend(current_date) and not is_included_date:
                    current_date += timedelta(days=1)
                    continue

                # 제외일인 경우 스킵
                if any(exclude_date and exclude_date.date() == day
                      for exclude_date in final_exclude_dates):
                    current_date += timedelta(days=1)
                    continue

                # 재고가 안전재고 이하가 되는지 확인
                if remaining_inventory <= safety_stock:
                    # 포함일(주말)인 경우 다음 월요일을 입고일로 설정
                    actual_purchase_date = (
                        get_next_monday(current_date) if is_included_date
                        else current_date
                    )

                    # 아직 도착하지 않은 입고 예정이 있는데 그 전에 안전재고에 도달한 경우
                    pending = [d for d in item_receipts if day < d <= end_date.date()]
                    if pending:
                        review_notes.append(
                            f"{actual_purchase_date:%Y-%m-%d} 발주 ← {min(pending):%Y-%m-%d} 입고예정 전 안전재고 도달"
                        )

                    safety_stock_info.append({
                        'reach_date': current_date.strftime('%Y-%m-%d'),
                        'purchase_date': actual_purchase_date.strftime('%Y-%m-%d'),
                        'stock_before': f"{remaining_inventory:.2f}",
                        'stock_after': f"{remaining_inventory + purchase_amount:.2f}"
                    })
                    remaining_inventory += purchase_amount

                # 재고 계산
                remaining_inventory -= daily_usage
                current_date += timedelta(days=1)

            receipt_text = ', '.join(applied_receipts) if applied_receipts else '-'
            if out_of_range:
                receipt_text += f" (계산기간 외 {out_of_range}건 미반영)"
            review_text = ' / '.join(review_notes) if review_notes else '-'

            # 결과 저장
            if safety_stock_info:
                results.append({
                    '품목코드': item_code,
                    '품목명': item_name,
                    '입고 예정 반영': receipt_text,
                    '안전재고 도달 일자': ', '.join(info['reach_date'] for info in safety_stock_info),
                    '도달시 재고량': ', '.join(info['stock_before'] for info in safety_stock_info),
                    '입고후 재고량': ', '.join(info['stock_after'] for info in safety_stock_info),
                    '입고 필요일': ', '.join(info['purchase_date'] for info in safety_stock_info),
                    '검토 필요': review_text,
                })
            else:
                results.append({
                    '품목코드': item_code,
                    '품목명': item_name,
                    '입고 예정 반영': receipt_text,
                    '안전재고 도달 일자': '계산기간 내 안전재고 미도달',
                    '도달시 재고량': '-',
                    '입고후 재고량': '-',
                    '입고 필요일': '입고 필요 없음',
                    '검토 필요': review_text,
                })

        return results, receipt_events

    except Exception as e:
        st.error(f"계산 중 에러 발생: {str(e)}")
        return None, []

def calculate_monthly_purchase(results_data, df):
    """월별 구매 필요량을 계산하는 함수"""
    monthly_data = []

    for _, row in df.iterrows():
        item_code = row['원료코드명']
        item_name = row['원료명']

        item_result = next((r for r in results_data if r['품목코드'] == item_code), None)

        if item_result and item_result['입고 필요일'] != '입고 필요 없음':
            purchase_dates = item_result['입고 필요일'].split(', ')
            stock_before = [float(x) for x in item_result['도달시 재고량'].split(', ')]
            stock_after = [float(x) for x in item_result['입고후 재고량'].split(', ')]

            actual_purchases = [after - before for before, after in zip(stock_before, stock_after)]

            monthly_counts = {}

            for date_str, purchase_amount in zip(purchase_dates, actual_purchases):
                date = datetime.strptime(date_str, '%Y-%m-%d')
                year = date.year
                month = date.month
                month_key = f"{year}-{month:02d}"
                display_key = f"{month}월({year}년)"

                if display_key not in monthly_counts:
                    monthly_counts[display_key] = {
                        'count': 0,
                        'amount': 0,
                        'sort_key': month_key
                    }
                monthly_counts[display_key]['count'] += 1
                monthly_counts[display_key]['amount'] += purchase_amount

            item_data = {
                '품목코드': item_code,
                '품목명': item_name
            }

            sorted_months = sorted(monthly_counts.items(),
                                 key=lambda x: x[1]['sort_key'])

            for month_display, data in sorted_months:
                item_data[f'{month_display} 입고 횟수'] = data['count']
                item_data[f'{month_display} 총 구매량'] = f"{data['amount']:.2f}"

            monthly_data.append(item_data)

    if not monthly_data:
        return pd.DataFrame(columns=['품목코드', '품목명'])

    return pd.DataFrame(monthly_data)

def build_order_info_base(df, prefill_df=None):
    """업로드된 재고 데이터의 품목 목록으로 발주 정보 입력용 기본 표를 생성하는 함수
    prefill_df(기준정보 파일)가 있으면 품목코드 기준으로 미리 채움
    """
    base = pd.DataFrame({
        '원료코드명': df['원료코드명'].map(normalize_code),
        '원료명': df['원료명'].astype(str),
    })
    base['부가세 구분'] = None
    base['제품 구분'] = None
    base['업체명'] = None
    base['단량(KG)'] = None
    base['단가(원)'] = None
    base['담당자 이메일'] = ''

    if prefill_df is not None:
        pre = prefill_df.copy()
        code_col = next((c for c in ['원료코드명', '품목코드'] if c in pre.columns), None)
        if code_col:
            pre[code_col] = pre[code_col].map(normalize_code)
            pre = pre.drop_duplicates(subset=[code_col], keep='first').set_index(code_col)

            for col in ORDER_INFO_COLUMNS:
                if col in pre.columns:
                    mapped = base['원료코드명'].map(pre[col])
                    if col in ['단량(KG)', '단가(원)']:
                        base[col] = pd.to_numeric(mapped, errors='coerce')
                    else:
                        cleaned = mapped.fillna('').astype(str).str.strip().replace('미입력', '')
                        base[col] = cleaned.where(cleaned != '', None)
        else:
            st.warning("⚠️ 기준정보 파일에 '품목코드' 또는 '원료코드명' 열이 없어 미리 채우기를 건너뜁니다.")

    return base

def create_purchase_order_df(results_data, df, order_info_df, vendor_map):
    """계산 결과 + 입력받은 발주 정보를 발주서 생성창 규격으로 변환하는 함수
    담당자 이메일이 비어 있으면 거래처 DB에서 자동 보완
    총액(원) = 입고량(KG) × 단가(원)  ※ 단가는 KG당 단가 기준
    """
    order_columns = ['입고일', '품목명', '부가세 구분', '제품 구분', '업체명',
                     '단량(KG)', '입고량(KG)', '단가(원)', '총액(원)', '담당자 이메일']

    info = order_info_df.copy()
    info['원료코드명'] = info['원료코드명'].map(normalize_code)
    info_lookup = info.set_index('원료코드명').to_dict('index')

    def clean(value):
        return '' if is_blank(value) else str(value).strip()

    order_rows = []

    for _, row in df.iterrows():
        item_code = normalize_code(row['원료코드명'])
        item_result = next((r for r in results_data if normalize_code(r['품목코드']) == item_code), None)
        if not item_result or item_result['입고 필요일'] == '입고 필요 없음':
            continue

        item_info = info_lookup.get(item_code, {})
        purchase_amount = float(row['1회 구매량'])  # 입고량(KG)
        unit_price = pd.to_numeric(item_info.get('단가(원)'), errors='coerce')
        total_price = round(purchase_amount * unit_price) if pd.notna(unit_price) else ''

        vendor = clean(item_info.get('업체명'))
        email = clean(item_info.get('담당자 이메일'))
        if not email and vendor:
            email = vendor_map.get(vendor, '')

        for date_str in item_result['입고 필요일'].split(', '):
            order_rows.append({
                '입고일': date_str,
                '품목명': row['원료명'],
                '부가세 구분': clean(item_info.get('부가세 구분')),
                '제품 구분': clean(item_info.get('제품 구분')),
                '업체명': vendor,
                '단량(KG)': '' if is_blank(item_info.get('단량(KG)')) else item_info.get('단량(KG)'),
                '입고량(KG)': purchase_amount,
                '단가(원)': unit_price if pd.notna(unit_price) else '',
                '총액(원)': total_price,
                '담당자 이메일': email,
            })

    if not order_rows:
        return pd.DataFrame(columns=order_columns)

    return pd.DataFrame(order_rows, columns=order_columns).sort_values(['입고일', '품목명']).reset_index(drop=True)

def create_calendar_sheet_safe(writer, results_data, receipt_events=None):
    """달력 형식으로 발주 필요일/입고 예정일을 표시하는 함수 (안전한 버전)"""
    try:
        order_dates, receipt_dates, all_dates = collect_calendar_dates(results_data, receipt_events)
        workbook = writer.book
        calendar_sheet = workbook.create_sheet('입고일정달력')

        if not all_dates:
            calendar_sheet.cell(row=1, column=1, value="입고 일정이 없습니다.")
            return

        start_date = min(all_dates)
        end_date = max(all_dates)

        current_date = start_date.replace(day=1)
        end_month = end_date.replace(day=1)

        header_fill = PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')
        weekend_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        border = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))

        current_row = 1

        while current_date <= end_month:
            month_title = current_date.strftime('%Y년 %m월')
            calendar_sheet.merge_cells(f'A{current_row}:G{current_row}')
            title_cell = calendar_sheet.cell(row=current_row, column=1, value=month_title)
            title_cell.font = Font(bold=True, size=12)
            title_cell.alignment = Alignment(horizontal='center')

            header_row = current_row + 1
            for col, day in enumerate(WEEKDAYS_KR, 1):
                cell = calendar_sheet.cell(row=header_row, column=col, value=day)
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
                calendar_sheet.column_dimensions[get_column_letter(col)].width = 15

            cal = calendar.monthcalendar(current_date.year, current_date.month)
            for week_idx, week in enumerate(cal):
                row = current_row + 2 + week_idx
                for day_idx, day in enumerate(week):
                    cell = calendar_sheet.cell(row=row, column=day_idx + 1)
                    cell.border = border
                    cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)

                    if day != 0:
                        date_str = f"{current_date.year}-{current_date.month:02d}-{day:02d}"
                        lines = [str(day)]
                        lines += order_dates.get(date_str, [])
                        lines += [f"[입고예정] {name}" for name in receipt_dates.get(date_str, [])]
                        cell.value = "\n".join(lines)

                        if day_idx >= 5:
                            cell.fill = weekend_fill

            current_row += len(cal) + 4
            current_date = (current_date + timedelta(days=32)).replace(day=1)

        for row_num in range(1, current_row):
            calendar_sheet.row_dimensions[row_num].height = 60

    except Exception as e:
        st.warning(f"달력 생성 중 오류: {str(e)}")
        workbook = writer.book
        if '입고일정달력' not in workbook.sheetnames:
            calendar_sheet = workbook.create_sheet('입고일정달력')
            calendar_sheet.cell(row=1, column=1, value=f"달력 생성 중 오류 발생: {str(e)}")

def create_excel_file(results_data, monthly_df, po_df, receipt_events=None):
    """엑셀 파일을 메모리에 생성하여 반환"""
    output = io.BytesIO()

    try:
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # 발주서 시트
            if not po_df.empty:
                po_df.to_excel(writer, sheet_name='발주서', index=False)
            else:
                pd.DataFrame({'메시지': ['계산 기간 내 발주가 필요한 품목이 없습니다.']}).to_excel(
                    writer, sheet_name='발주서', index=False)

            # 입고계획 시트
            pd.DataFrame(results_data).to_excel(writer, sheet_name='입고계획', index=False)

            # 입고예정 시트 (입력한 경우만)
            if receipt_events:
                pd.DataFrame(receipt_events).sort_values('날짜').to_excel(
                    writer, sheet_name='입고예정', index=False)

            # 월별구매량 시트
            if not monthly_df.empty:
                monthly_df.to_excel(writer, sheet_name='월별구매량', index=False)
            else:
                pd.DataFrame({'메시지': ['계산 기간 내 구매가 필요한 품목이 없습니다.']}).to_excel(
                    writer, sheet_name='월별구매량', index=False)

            create_calendar_sheet_safe(writer, results_data, receipt_events)

        output.seek(0)
        return output
    except Exception as e:
        st.error(f"Excel 파일 생성 중 오류: {str(e)}")
        return None

def render_vendor_db_section(vendor_db):
    """거래처 DB 확인·업로드(전체 교체)·다운로드 화면"""
    with st.expander(f"🏢 거래처 DB — 현재 {len(vendor_db)}개 등록", expanded=vendor_db.empty):
        if st.session_state.pop('vendor_saved_msg', None):
            st.success("✅ 거래처 DB가 업로드한 파일로 교체되었습니다. (이전 DB는 거래처DB_backup.csv 로 백업)")

        st.caption("거래처가 추가·변경되면 '업체명', '담당자 이메일' 열이 있는 Excel/CSV 파일을 업로드하세요. "
                   "업로드한 파일 내용으로 DB 전체가 교체됩니다. 이메일 '미입력'은 빈 값으로 처리됩니다.")

        if vendor_db.empty:
            st.info("등록된 거래처가 없습니다. 거래처 파일을 업로드하세요.")
        else:
            st.dataframe(vendor_db, width='stretch', height=250)

        vendor_file = st.file_uploader(
            "거래처 파일 업로드 (전체 교체)",
            type=['xlsx', 'xls', 'csv'],
            key="vendor_upload",
        )

        if vendor_file is not None:
            try:
                new_vdb = normalize_vendor_df(read_table_file(vendor_file, vendor_file.name))

                old_map = dict(zip(vendor_db['업체명'], vendor_db['담당자 이메일']))
                new_map = dict(zip(new_vdb['업체명'], new_vdb['담당자 이메일']))
                added = [v for v in new_map if v not in old_map]
                removed = [v for v in old_map if v not in new_map]
                changed = [v for v in new_map if v in old_map and old_map[v] != new_map[v]]

                if not (added or removed or changed):
                    st.info("업로드한 파일이 현재 거래처 DB와 동일합니다.")
                else:
                    st.write(f"**업로드 파일: {len(new_vdb)}개 업체** — "
                             f"신규 {len(added)}개 / 삭제 {len(removed)}개 / 이메일 변경 {len(changed)}개")
                    if added:
                        st.caption("신규: " + ", ".join(added))
                    if changed:
                        st.caption("이메일 변경: " + ", ".join(changed))
                    if removed:
                        st.warning("삭제될 업체: " + ", ".join(removed))

                    if st.button("✅ 이 파일로 거래처 DB 교체", type="primary", width='stretch'):
                        save_vendor_db(new_vdb)
                        st.session_state['vendor_saved_msg'] = True
                        st.rerun()
            except Exception as e:
                st.error(f"거래처 파일을 읽지 못했습니다: {str(e)}")

        if not vendor_db.empty:
            buffer = io.BytesIO()
            vendor_db.to_excel(buffer, index=False, sheet_name='거래처DB')
            buffer.seek(0)
            st.download_button(
                "📥 현재 거래처 DB 다운로드 (수정 후 다시 업로드)",
                data=buffer,
                file_name="거래처DB.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width='stretch',
            )

# Streamlit 앱 시작
def main():
    st.title("📦 재고 관리 시스템")
    st.markdown("---")

    # 거래처 DB 로드
    vendor_db = load_vendor_db()
    vendor_map = {v: e for v, e in zip(vendor_db['업체명'], vendor_db['담당자 이메일']) if e}

    # 사이드바
    st.sidebar.header("📝 사용 방법")
    st.sidebar.markdown("""
    1. **Excel 파일 업로드**: 재고 데이터가 포함된 Excel 파일을 업로드하세요.
    2. **발주 정보 입력**: 품목별 구분·업체·단가 등을 입력하세요. 업체 선택 시 이메일이 자동 입력됩니다.
    3. **입고 예정 입력**: 이미 발주해서 들어올 물량이 있으면 일자·수량을 입력하세요.
    4. **계산 실행**: '계산 시작' 버튼을 클릭하세요.
    5. **결과 확인**: 발주서·입고 계획·월별 구매량을 확인하세요.
    6. **복사/다운로드**: 발주서 탭에서 복사하거나 Excel로 다운로드하세요.
    """)

    st.sidebar.markdown("---")
    st.sidebar.header("📋 필요한 열 정보 (업로드 파일)")
    st.sidebar.markdown("""
    Excel 파일에는 다음 열이 포함되어야 합니다:
    - 원료코드명
    - 원료명
    - 현재 재고
    - 최소 사용량
    - 최대 사용량
    - 계산 일자
    - 안전 재고
    - 1회 구매량
    - 제외 날짜 (선택사항)
    - 포함 날짜 (선택사항)
    """)

    # 거래처 DB 관리 (업로드로 교체)
    render_vendor_db_section(vendor_db)

    # 파일 업로드
    uploaded_file = st.file_uploader(
        "Excel 파일을 업로드하세요",
        type=['xlsx', 'xls'],
        help="재고 데이터가 포함된 Excel 파일을 선택하세요."
    )

    if uploaded_file is not None:
        try:
            df = pd.read_excel(uploaded_file)
            today = datetime.now().date()

            # 다른 파일이 업로드되면 이전 계산 결과 초기화
            if st.session_state.get('results_file') != uploaded_file.name:
                st.session_state.pop('results', None)

            # 데이터 미리보기
            st.subheader("📊 업로드된 데이터 미리보기")
            st.dataframe(df.head(), width='stretch')

            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("총 품목 수", len(df))
            with col2:
                st.metric("총 열 수", len(df.columns))
            with col3:
                st.metric("데이터 크기", f"{df.shape[0]} × {df.shape[1]}")

            # 필수 열 확인 (기존 업로드 형식 그대로)
            required_columns = [
                '원료코드명', '원료명', '현재 재고', '최소 사용량',
                '최대 사용량', '계산 일자', '안전 재고', '1회 구매량'
            ]

            missing_columns = [col for col in required_columns if col not in df.columns]

            if missing_columns:
                st.error(f"❌ 필수 열이 누락되었습니다: {', '.join(missing_columns)}")
                st.stop()
            else:
                st.success("✅ 모든 필수 열이 확인되었습니다!")

            # ─────────────────────────────────────────────
            # 발주 정보 입력 (계산 시작 전 단계)
            # ─────────────────────────────────────────────
            st.markdown("---")
            st.subheader("🧾 발주 정보 입력")
            st.caption("품목별 부가세 구분·제품 구분·업체명·단량·단가를 입력하세요. "
                       "업체명을 선택하면 담당자 이메일이 거래처 DB에서 자동 입력되며, 직접 수정도 가능합니다. "
                       "총액은 입고량 × 단가로 자동 계산됩니다.")

            # 기준정보 파일로 자동 채우기 (선택사항)
            prefill_file = st.file_uploader(
                "기준정보 파일로 자동 채우기 (선택사항)",
                type=['xlsx', 'xls'],
                key="prefill_file",
                help="'품목코드(또는 원료코드명)' 열과 업체명·단가 등의 열이 있으면 품목코드 기준으로 자동으로 채워집니다."
            )

            prefill_df = None
            if prefill_file is not None:
                try:
                    prefill_df = pd.read_excel(prefill_file)
                except Exception as e:
                    st.warning(f"⚠️ 기준정보 파일을 읽지 못했습니다: {str(e)}")

            # 발주 정보 표 초기화 (재고 파일/기준정보 파일이 바뀔 때만 새로 생성)
            editor_key = f"{uploaded_file.name}|{prefill_file.name if prefill_file else 'none'}"
            if st.session_state.get('order_info_key') != editor_key:
                st.session_state['order_info_base'] = build_order_info_base(df, prefill_df)
                st.session_state['order_info_key'] = editor_key
                st.session_state['vendor_autofill_map'] = {}

            # 업체명 드롭다운 옵션: 거래처 DB + 이미 입력된 값
            existing_vendors = {
                str(v).strip() for v in st.session_state['order_info_base']['업체명'].dropna()
                if str(v).strip()
            }
            vendor_options = sorted(set(vendor_db['업체명']) | existing_vendors)

            order_info_df = st.data_editor(
                st.session_state['order_info_base'],
                width='stretch',
                num_rows="fixed",
                disabled=['원료코드명', '원료명'],
                column_config={
                    '원료코드명': st.column_config.TextColumn('원료코드명'),
                    '원료명': st.column_config.TextColumn('원료명'),
                    '부가세 구분': st.column_config.SelectboxColumn('부가세 구분', options=VAT_OPTIONS, required=False),
                    '제품 구분': st.column_config.SelectboxColumn('제품 구분', options=PRODUCT_TYPE_OPTIONS, required=False),
                    '업체명': st.column_config.SelectboxColumn(
                        '업체명', options=vendor_options, required=False,
                        help="목록에 없는 업체는 위 '거래처 DB'에 파일을 업로드해 추가하세요."),
                    '단량(KG)': st.column_config.NumberColumn('단량(KG)', min_value=0, format="%.2f"),
                    '단가(원)': st.column_config.NumberColumn('단가(원)', min_value=0, format="%.0f"),
                    '담당자 이메일': st.column_config.TextColumn('담당자 이메일', help="업체 선택 시 자동 입력, 직접 수정 가능"),
                },
                key=f"order_editor_{editor_key}"
            )

            # 업체명 선택/변경 시 담당자 이메일 자동 입력 (수동 수정은 유지)
            autofill_map = st.session_state.setdefault('vendor_autofill_map', {})
            autofill_changed = False
            for idx in order_info_df.index:
                vendor_value = order_info_df.at[idx, '업체명']
                vendor_value = '' if is_blank(vendor_value) else str(vendor_value).strip()
                if vendor_value and vendor_value != autofill_map.get(idx):
                    autofill_map[idx] = vendor_value
                    db_email = vendor_map.get(vendor_value, '')
                    if db_email:
                        order_info_df.at[idx, '담당자 이메일'] = db_email
                        autofill_changed = True

            # 편집 내용을 세션에 보관 (계산 후에도 유지)
            st.session_state['order_info_base'] = order_info_df
            if autofill_changed:
                st.rerun()

            filled_price = int(pd.to_numeric(order_info_df['단가(원)'], errors='coerce').notna().sum())
            filled_vendor = int(order_info_df['업체명'].fillna('').astype(str).str.strip().ne('').sum())
            st.caption(f"입력 현황 — 단가: {filled_price}/{len(order_info_df)}개, 업체명: {filled_vendor}/{len(order_info_df)}개 "
                       "(비워두면 발주서에서 해당 칸이 빈 값으로 출력됩니다)")

            # ─────────────────────────────────────────────
            # 입고 예정 입력 (계산에 반영)
            # ─────────────────────────────────────────────
            st.markdown("---")
            st.subheader("📥 입고 예정 입력 (선택사항)")
            st.caption("이미 발주해서 들어올 예정인 물량을 입력하면, 해당 일자 시작 시점에 재고로 더해서 계산합니다. "
                       "오늘 이미 입고되어 '현재 재고'에 포함된 물량은 입력하지 마세요. "
                       "행 추가는 표 아래 ➕, Excel에서 복사한 여러 행을 붙여넣기도 가능합니다.")

            item_codes = df['원료코드명'].map(normalize_code)
            item_labels = [f"{code} | {name}" for code, name in zip(item_codes, df['원료명'].astype(str))]
            label_to_code = dict(zip(item_labels, item_codes))

            receipt_base_key = f"receipt_base_{uploaded_file.name}"
            if receipt_base_key not in st.session_state:
                st.session_state[receipt_base_key] = pd.DataFrame({
                    '원료': pd.Series(dtype='object'),
                    '입고일': pd.Series(dtype='datetime64[ns]'),
                    '입고량(KG)': pd.Series(dtype='float'),
                })

            # 입력 원본(base)은 고정하고 편집 결과만 읽음 (동적 행 중복 방지)
            receipts_edited = st.data_editor(
                st.session_state[receipt_base_key],
                num_rows="dynamic",
                width='stretch',
                key=f"receipt_editor_{uploaded_file.name}",
                column_config={
                    '원료': st.column_config.SelectboxColumn('원료', options=item_labels, width='large'),
                    '입고일': st.column_config.DateColumn('입고일', min_value=today, format="YYYY-MM-DD"),
                    '입고량(KG)': st.column_config.NumberColumn('입고량(KG)', min_value=0, format="%.2f"),
                },
            )

            scheduled_receipts, past_count, incomplete_count = parse_scheduled_receipts(
                receipts_edited, label_to_code, today)

            receipt_count = sum(len(v) for v in scheduled_receipts.values())
            receipt_total = sum(q for v in scheduled_receipts.values() for q in v.values())
            if receipt_count:
                st.caption(f"계산에 반영될 입고 예정: {receipt_count}건, 총 {receipt_total:,.2f}KG "
                           "(같은 원료·같은 날짜는 합산)")
            if incomplete_count:
                st.warning(f"⚠️ 원료·입고일·입고량 중 빠진 값이 있는 {incomplete_count}개 행은 계산에서 제외됩니다.")
            if past_count:
                st.warning(f"⚠️ 오늘 이전 날짜 {past_count}건은 이미 현재 재고에 포함된 것으로 보고 제외합니다.")

            # ─────────────────────────────────────────────
            # 날짜 설정
            # ─────────────────────────────────────────────
            st.markdown("---")
            st.subheader("📅 날짜 설정 (선택사항)")

            if 'exclude_dates_list' not in st.session_state:
                st.session_state.exclude_dates_list = []
            if 'include_dates_list' not in st.session_state:
                st.session_state.include_dates_list = []

            date_col1, date_col2 = st.columns(2)
            max_date = today + timedelta(days=365)

            with date_col1:
                st.markdown("**🚫 제외할 날짜들**")
                st.caption("계산에서 제외할 날짜를 추가하세요 (휴일, 비가동일 등)")

                exclude_date_input = st.date_input(
                    "제외할 날짜 선택",
                    value=today,
                    min_value=today,
                    max_value=max_date,
                    key="exclude_date_input"
                )

                col1, col2 = st.columns(2)
                with col1:
                    if st.button("➕ 제외 날짜 추가", width='stretch'):
                        if exclude_date_input not in st.session_state.exclude_dates_list:
                            st.session_state.exclude_dates_list.append(exclude_date_input)
                            st.success(f"제외 날짜 추가: {exclude_date_input.strftime('%Y-%m-%d')}")
                        else:
                            st.warning("이미 추가된 날짜입니다.")

                with col2:
                    if st.button("🗑️ 전체 삭제", width='stretch'):
                        st.session_state.exclude_dates_list = []
                        st.success("모든 제외 날짜가 삭제되었습니다.")

                if st.session_state.exclude_dates_list:
                    st.write("**추가된 제외 날짜들:**")
                    for i, date in enumerate(sorted(st.session_state.exclude_dates_list)):
                        col_date, col_del = st.columns([3, 1])
                        with col_date:
                            weekday = WEEKDAYS_KR[date.weekday()]
                            st.write(f"• {date.strftime('%Y-%m-%d')} ({weekday})")
                        with col_del:
                            if st.button("❌", key=f"del_exclude_{i}", help="이 날짜 삭제"):
                                st.session_state.exclude_dates_list.remove(date)
                                st.rerun()
                else:
                    st.info("추가된 제외 날짜가 없습니다.")

            with date_col2:
                st.markdown("**✅ 포함할 날짜들**")
                st.caption("주말이지만 가동하는 날짜를 추가하세요")

                include_date_input = st.date_input(
                    "포함할 날짜 선택",
                    value=today,
                    min_value=today,
                    max_value=max_date,
                    key="include_date_input"
                )

                col1, col2 = st.columns(2)
                with col1:
                    if st.button("➕ 포함 날짜 추가", width='stretch'):
                        if include_date_input not in st.session_state.include_dates_list:
                            st.session_state.include_dates_list.append(include_date_input)
                            weekday = WEEKDAYS_KR[include_date_input.weekday()]
                            st.success(f"포함 날짜 추가: {include_date_input.strftime('%Y-%m-%d')} ({weekday})")
                        else:
                            st.warning("이미 추가된 날짜입니다.")

                with col2:
                    if st.button("🗑️ 전체 삭제 ", width='stretch', key="clear_include"):
                        st.session_state.include_dates_list = []
                        st.success("모든 포함 날짜가 삭제되었습니다.")

                if st.session_state.include_dates_list:
                    st.write("**추가된 포함 날짜들:**")
                    for i, date in enumerate(sorted(st.session_state.include_dates_list)):
                        col_date, col_del = st.columns([3, 1])
                        with col_date:
                            weekday = WEEKDAYS_KR[date.weekday()]
                            weekend_text = " 🟡" if date.weekday() >= 5 else " 🔵"
                            st.write(f"• {date.strftime('%Y-%m-%d')} ({weekday}){weekend_text}")
                        with col_del:
                            if st.button("❌", key=f"del_include_{i}", help="이 날짜 삭제"):
                                st.session_state.include_dates_list.remove(date)
                                st.rerun()
                else:
                    st.info("추가된 포함 날짜가 없습니다.")

            total_dates = len(st.session_state.exclude_dates_list) + len(st.session_state.include_dates_list)
            if total_dates > 0:
                st.success(f"📝 총 {total_dates}개의 날짜가 설정되었습니다. (제외: {len(st.session_state.exclude_dates_list)}개, 포함: {len(st.session_state.include_dates_list)}개)")
                st.caption("🔵 = 평일, 🟡 = 주말")

            # 계산 입력값 서명 (계산 후 입력이 바뀌었는지 확인용)
            calc_signature = str((
                sorted((code, sorted(dates.items())) for code, dates in scheduled_receipts.items()),
                sorted(st.session_state.exclude_dates_list),
                sorted(st.session_state.include_dates_list),
            ))

            if st.button("🚀 계산 시작", type="primary", width='stretch'):
                with st.spinner("계산 중입니다... 잠시만 기다려주세요."):
                    results, receipt_events = calculate_purchase_date(
                        df,
                        st.session_state.exclude_dates_list,
                        st.session_state.include_dates_list,
                        scheduled_receipts,
                    )
                    if results:
                        st.session_state['results'] = results
                        st.session_state['receipt_events'] = receipt_events
                        st.session_state['results_file'] = uploaded_file.name
                        st.session_state['calc_signature'] = calc_signature
                    else:
                        st.session_state.pop('results', None)
                        st.error("❌ 계산 중 오류가 발생했습니다.")

            # 계산 결과 표시
            results = st.session_state.get('results')
            if results:
                receipt_events = st.session_state.get('receipt_events', [])
                st.success("✅ 계산이 완료되었습니다!")
                if st.session_state.get('calc_signature') != calc_signature:
                    st.warning("⚠️ 계산 이후 입고 예정 또는 날짜 설정이 변경되었습니다. "
                               "'계산 시작'을 다시 눌러야 결과에 반영됩니다.")

                # 발주서/월별 데이터 준비 (발주 정보는 최신 입력값 반영)
                po_df = create_purchase_order_df(results, df, order_info_df, vendor_map)
                monthly_df = calculate_monthly_purchase(results, df)

                tab1, tab2, tab3, tab4, tab5 = st.tabs(
                    ["🧾 발주서", "📋 입고 계획", "📊 월별 구매량", "📅 입고 일정 요약", "🗓️ 캘린더 보기"])

                with tab1:
                    st.subheader("🧾 발주서 출력 (생성창 규격)")
                    st.caption("자동 산출된 발주 필요 건만 출력됩니다. 입력한 입고 예정 건은 이미 발주된 것으로 보고 제외합니다.")

                    if po_df.empty:
                        st.info("계산 기간 내 발주가 필요한 품목이 없습니다.")
                    else:
                        st.dataframe(po_df, width='stretch')

                        missing_price = int(po_df['총액(원)'].astype(str).eq('').sum())
                        if missing_price:
                            st.warning(f"⚠️ 단가 미입력으로 총액을 계산하지 못한 행이 {missing_price}건 있습니다. "
                                       "위 '발주 정보 입력' 표에서 단가를 채우면 즉시 반영됩니다.")

                        total_sum = pd.to_numeric(po_df['총액(원)'], errors='coerce').sum()
                        sum_col1, sum_col2 = st.columns(2)
                        with sum_col1:
                            st.metric("발주 건수", len(po_df))
                        with sum_col2:
                            st.metric("발주 총액(원)", f"{total_sum:,.0f}")

                        st.markdown("**📋 아래 박스 우측 상단의 복사 버튼을 눌러 발주서 생성창에 그대로 붙여넣으세요.** "
                                    "(탭 구분 형식이라 Excel/그리드에 열이 맞춰 들어갑니다)")
                        include_header = st.checkbox("헤더(열 이름) 포함", value=True, key="po_header")
                        tsv_text = po_df.to_csv(sep='\t', index=False, header=include_header)
                        st.code(tsv_text, language=None)

                with tab2:
                    st.subheader("📋 입고 계획")
                    output_df = pd.DataFrame(results)
                    st.dataframe(output_df, width='stretch')

                    review_items = [r for r in results if r['검토 필요'] != '-']
                    if review_items:
                        with st.expander(f"⚠️ 검토 필요 {len(review_items)}개 품목 — 입고 예정 전에 안전재고 도달", expanded=True):
                            st.caption("입고 예정 물량이 도착하기 전에 안전재고에 닿아 추가 발주가 산출된 품목입니다. "
                                       "예정 입고를 그 날짜 이전으로 당길 수 있다면 추가 발주가 필요 없을 수 있습니다.")
                            for r in review_items:
                                st.write(f"• **{r['품목명']}** ({r['품목코드']}): {r['검토 필요']}")

                    if st.session_state.exclude_dates_list or st.session_state.include_dates_list:
                        with st.expander("📅 적용된 날짜 설정 확인"):
                            col1, col2 = st.columns(2)
                            with col1:
                                if st.session_state.exclude_dates_list:
                                    st.write("**🚫 제외된 날짜들:**")
                                    for date in sorted(st.session_state.exclude_dates_list):
                                        weekday = WEEKDAYS_KR[date.weekday()]
                                        st.write(f"• {date.strftime('%Y-%m-%d')} ({weekday})")
                                else:
                                    st.write("**🚫 제외된 날짜:** 없음")

                            with col2:
                                if st.session_state.include_dates_list:
                                    st.write("**✅ 포함된 날짜들:**")
                                    for date in sorted(st.session_state.include_dates_list):
                                        weekday = WEEKDAYS_KR[date.weekday()]
                                        weekend_text = " 🟡" if date.weekday() >= 5 else " 🔵"
                                        st.write(f"• {date.strftime('%Y-%m-%d')} ({weekday}){weekend_text}")
                                else:
                                    st.write("**✅ 포함된 날짜:** 없음")

                    purchase_needed = len([r for r in results if r['입고 필요일'] != '입고 필요 없음'])
                    no_purchase = len(results) - purchase_needed

                    col1, col2, col3 = st.columns(3)
                    with col1:
                        st.metric("입고 필요 품목", purchase_needed)
                    with col2:
                        st.metric("입고 불필요 품목", no_purchase)
                    with col3:
                        if len(results) > 0:
                            ratio = (purchase_needed / len(results)) * 100
                            st.metric("입고 필요 비율", f"{ratio:.1f}%")

                with tab3:
                    st.subheader("📊 월별 구매량")
                    if not monthly_df.empty:
                        st.dataframe(monthly_df, width='stretch')
                    else:
                        st.info("계산 기간 내 구매가 필요한 품목이 없습니다.")

                with tab4:
                    st.subheader("📅 입고 일정 요약")

                    schedule_rows = []
                    for item in results:
                        if item['입고 필요일'] != '입고 필요 없음':
                            for date in item['입고 필요일'].split(', '):
                                schedule_rows.append({
                                    '날짜': date,
                                    '구분': '발주 필요',
                                    '품목명': item['품목명'],
                                    '품목코드': item['품목코드'],
                                })
                    for event in receipt_events:
                        schedule_rows.append({
                            '날짜': event['날짜'],
                            '구분': '입고 예정',
                            '품목명': event['품목명'],
                            '품목코드': event['품목코드'],
                        })

                    if schedule_rows:
                        schedule_df = pd.DataFrame(schedule_rows)
                        schedule_df['날짜'] = pd.to_datetime(schedule_df['날짜'])
                        schedule_df = schedule_df.sort_values(['날짜', '구분'])
                        schedule_df['날짜'] = schedule_df['날짜'].dt.strftime('%Y-%m-%d')

                        st.dataframe(schedule_df, width='stretch')

                        date_counts = schedule_df.groupby(['날짜', '구분']).size().unstack(fill_value=0)
                        st.bar_chart(date_counts)
                    else:
                        st.info("계산 기간 내 입고 일정이 없습니다.")

                with tab5:
                    st.subheader("🗓️ 입고 일정 캘린더")
                    display_calendar(results, receipt_events)

                # 다운로드 버튼
                st.markdown("---")
                st.subheader("💾 결과 다운로드")

                try:
                    excel_file = create_excel_file(results, monthly_df, po_df, receipt_events)

                    if excel_file is not None:
                        st.download_button(
                            label="📥 Excel 파일 다운로드 (발주서 포함)",
                            data=excel_file,
                            file_name=f"입고계획_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            width='stretch'
                        )
                    else:
                        st.error("Excel 파일 생성에 실패했습니다.")
                except Exception as e:
                    st.error(f"다운로드 파일 생성 중 오류: {str(e)}")
                    csv = po_df.to_csv(index=False, encoding='utf-8-sig')
                    st.download_button(
                        label="📄 CSV 파일 다운로드 (발주서)",
                        data=csv,
                        file_name=f"발주서_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                        mime="text/csv",
                        width='stretch'
                    )

        except Exception as e:
            st.error(f"❌ 파일 처리 중 오류가 발생했습니다: {str(e)}")

    else:
        st.info("👆 Excel 파일을 업로드해주세요.")

        st.subheader("📋 샘플 데이터 형식")
        sample_data = {
            '원료코드명': ['1010101', '1010111'],
            '원료명': ['닭고기 MDCM', '닭고기가수분해단백질분말'],
            '현재 재고': [100, 50],
            '최소 사용량': [5, 3],
            '최대 사용량': [10, 8],
            '계산 일자': [30, 30],
            '안전 재고': [20, 15],
            '1회 구매량': [100, 80],
            '제외 날짜': ['2026-10-09', ''],
            '포함 날짜': ['', '2026-10-17']
        }
        st.dataframe(pd.DataFrame(sample_data), width='stretch')

if __name__ == "__main__":
    main()
